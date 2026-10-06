"""Read the catalogue workbook into plain Python values."""

import datetime as dt
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from oercat.schema import SHEET_DATA, SHEET_PICKLISTS


@dataclass
class RawRow:
    excel_row: int  # 1-based row number as shown in Excel
    cells: dict[str, str]  # stripped header -> cell text ("" if empty)


@dataclass
class Workbook:
    headers: list[str]
    rows: list[RawRow]
    picklists: dict[str, list[str]]  # attribute name -> allowed values, in sheet order
    modified: dt.datetime | None = None  # "last saved" timestamp from the file properties


def cell_text(value) -> str:
    """Convert an openpyxl cell value to text the way it appears in Excel."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value).strip()


def read_workbook(path: Path) -> Workbook:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET_DATA] if SHEET_DATA in wb.sheetnames else wb.worksheets[0]

    rows_iter = ws.iter_rows(values_only=True)
    headers = [cell_text(h) for h in next(rows_iter)]
    rows = []
    for excel_row, values in enumerate(rows_iter, start=2):
        cells = {h: cell_text(v) for h, v in zip(headers, values) if h}
        if any(cells.values()):
            rows.append(RawRow(excel_row, cells))

    picklists: dict[str, list[str]] = {}
    if SHEET_PICKLISTS in wb.sheetnames:
        it = wb[SHEET_PICKLISTS].iter_rows(values_only=True)
        next(it)  # header
        for attribute, value, _desc, _iri, kind, *_ in it:
            if cell_text(kind) != "CVL" and attribute and value:  # PI (one row lacks the marker)
                picklists.setdefault(cell_text(attribute), []).append(cell_text(value))

    modified = wb.properties.modified
    wb.close()
    return Workbook([h for h in headers if h], rows, picklists, modified)
